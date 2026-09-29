#!/usr/bin/env python3
"""Extract a standardized human-reference table from the complete human
annotation Excel (Conversation sheet, schema with Trial_word / Intended_word /
Type MP-NMP-F / Word_Start), for use by build_comparison_excel.py and other
comparison tools.

Keeps minimal-pair rows only (Type == MP) with a valid time, maps each word to
its pair via master_vocab.csv, and computes Bark-normalized formants with the
same formula as the pipeline: bark = 26.81/(1+1960/F) - 0.53.

Output columns: Pair, Speaker, Word, Trial_word, Intended_word, Perceived_word,
Onset, F1, F2, F1_Bark, F2_Bark

Usage: python3 extract_human_reference.py COMPLETE_HUMAN.xlsx master_vocab.csv OUT.xlsx
"""
import sys, os, csv, re
import openpyxl
from openpyxl.styles import Font

clean = lambda s: re.sub(r"[^a-z']", "", str(s).lower().strip())
bark = lambda F: 26.81 / (1 + 1960 / F) - 0.53

def _to_number(v):
    try:
        f = float(v)
        # NaN check
        return f if f == f else None
    except (TypeError, ValueError):
        return None

if len(sys.argv) != 4:
    sys.exit("Usage: python3 extract_human_reference.py "
             "COMPLETE_HUMAN.xlsx master_vocab.csv OUT.xlsx")
src, vocab_csv, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
for path in (src, vocab_csv):
    if not os.path.isfile(path):
        sys.exit(f"ERROR: '{path}' not found.")

# word -> pair from the master vocabulary
word2pair = {}
with open(vocab_csv, encoding="utf-8") as f:
    for row in csv.DictReader(f):
        w, p = clean(row["word"]), str(row.get("pair") or "").strip()
        if w and p:
            word2pair[w] = p

def _normalise_header(h):
    """Tidies a column heading so it can be matched regardless of capitals
    or spacing."""
    return re.sub(r"\s+", "", str(h or "").lower())

wb = openpyxl.load_workbook(src, data_only=True)
# find the conversation sheet by content, not by name: the one whose header
# row contains a Trial_word-like column
ws = None
for sheet in wb.worksheets:
    if any("trial_word" in _normalise_header(c.value) for c in sheet[1]):
        ws = sheet
        break
if ws is None:
    sys.exit(f"ERROR: none of the sheets in '{src.split('/')[-1]}' looks like "
             "the conversation annotation (no header containing 'Trial_word'). "
             f"Sheets found: {', '.join(wb.sheetnames)}. Expected: the same "
             "expected format (a sheet with columns "
             "like Type(MP/NMP/F), Trial_word, Intended_word, "
             "converted time / Word_Start, F1, F2).")

hdr = [str(c.value or "").strip() for c in ws[1]]
nhdr = [_normalise_header(h) for h in hdr]
def _find_column(*keys):
    for k in keys:
        for i, h in enumerate(nhdr):
            if _normalise_header(k) in h:
                return i
    return None

iType  = _find_column("type(mp", "type (mp")
# the column holding the two speakers, whatever the annotation calls them
iSpk   = _find_column("participant", "speaker", "talker")
iConv  = _find_column("converted time", "convertedtime", "normalized time")
iWS    = _find_column("word_start")
iTrial = _find_column("trial_word")
iInt   = _find_column("intended_word")
iPerc  = _find_column("perceived_word")
iF1    = nhdr.index("f1") if "f1" in nhdr else None
iF2    = nhdr.index("f2") if "f2" in nhdr else None

missing = [name for name, idx in
           [("Type(MP/NMP/F)", iType), ("a speaker column (Participant, Speaker or Talker)", iSpk),
            ("a time column (converted time or Word_Start)",
             iWS if iWS is not None else iConv)]
           if idx is None]
if missing:
    sys.exit(f"ERROR: sheet '{ws.title}' of '{src.split('/')[-1]}' is missing "
             f"required column(s): {', '.join(missing)}. Columns found: "
             f"{', '.join(h for h in hdr if h)}")

rows_out, skipped = [], 0
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(r, c + 1).value for c in range(len(hdr))]
    if str(vals[iType] or "").strip().upper() != "MP":
        continue
    onset = ((_to_number(vals[iWS]) if iWS is not None else None)
             or (_to_number(vals[iConv]) if iConv is not None else None))
    if onset is None:
        skipped += 1
        continue
    # word: intended first (most consistent with the acoustics, per the
    # annotation guide), then trial word, then perceived
    word = next((str(vals[i]).strip() for i in (iInt, iTrial, iPerc)
                 if i is not None and vals[i] not in (None, "")), "")
    pair = next((word2pair[clean(str(vals[i]))] for i in (iInt, iTrial, iPerc)
                 if i is not None and vals[i] is not None
                 and clean(str(vals[i])) in word2pair), "")
    if not pair:
        skipped += 1
        continue
    f1, f2 = _to_number(vals[iF1]) if iF1 is not None else None, _to_number(vals[iF2]) if iF2 is not None else None
    # normalize speaker capitalization ("participant" -> "Participant")
    rows_out.append([pair, str(vals[iSpk] or "").strip().capitalize(), word,
                     str(vals[iTrial] or ""), str(vals[iInt] or ""), str(vals[iPerc] or ""),
                     round(onset, 3),
                     f1, f2,
                     round(bark(f1), 4) if f1 else None,
                     round(bark(f2), 4) if f2 else None])

out = openpyxl.Workbook()
o = out.active; o.title = "HumanRef"
o.append(["Pair", "Speaker", "Word", "Trial_word", "Intended_word", "Perceived_word",
          "Onset", "F1", "F2", "F1_Bark", "F2_Bark"])
for c in o[1]: c.font = Font(bold=True)
for row in rows_out:
    o.append(row)
out.save(out_path)
print(f"MP tokens extracted: {len(rows_out)} (skipped {skipped}: no time or unknown word)")
print(f"-> {out_path}")

"""Reads word timestamps from an annotation Excel.

The pipeline's first three stages exist only to produce a list of words with
their times and their speaker. Where that list already exists, this module
reads it instead and the measurement carries on from stage 4 with the same
code, so a measurement made this way and one made from the recordings are
directly comparable.
"""
import os

from openpyxl import load_workbook

# how long a word is taken to be when the annotation gives only its onset;
# the boundaries are then refined by silence detection inside that window
ONSET_ONLY_WINDOW = 0.6

_WORD = ("word", "trial_word", "intended_word", "target", "token")
_ONSET = ("onset", "word_start", "start", "tstart", "begin", "start_time")
_OFFSET = ("offset", "word_end", "end", "tend", "stop", "end_time")
_SPEAKER = ("speaker", "participant", "talker")


def _key(cell):
    return str(cell or "").strip().lower().replace(" ", "_")


def _column(header, names):
    for n in names:
        if n in header:
            return header[n]
    return None


def _number(v):
    if v is None or str(v).strip() == "":
        return None
    try:
        return float(str(v).strip().replace(",", "."))
    except ValueError:
        return None


def read_annotations(path, audio_by_speaker, default_speaker):
    """Rows ready for measurement, from the first usable sheet of the Excel.

    A word column and an onset column are required; an end column is used
    where it is there, and where it is not the row is given a short window
    from its onset and marked so the output says which rows were treated
    that way.
    """
    if not os.path.isfile(path):
        raise FileNotFoundError(f"annotation file not found: {path}")
    wb = load_workbook(path, data_only=True)

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = {}
        for i, cell in enumerate(rows[0]):
            header.setdefault(_key(cell), i)
        i_word = _column(header, _WORD)
        i_on = _column(header, _ONSET)
        if i_word is None or i_on is None:
            continue
        i_off = _column(header, _OFFSET)
        i_spk = _column(header, _SPEAKER)
        break
    else:
        raise RuntimeError(
            f"{os.path.basename(path)} has no sheet with both a word column "
            f"and an onset column. Expected one of {', '.join(_WORD)} for the "
            f"word, and one of {', '.join(_ONSET)} for the onset.")

    known = {str(k).strip().lower(): k for k in audio_by_speaker}
    out, skipped, unknown = [], 0, set()
    for raw in rows[1:]:
        if not raw or i_word >= len(raw):
            continue
        word = str(raw[i_word] or "").strip()
        onset = _number(raw[i_on]) if i_on < len(raw) else None
        if not word or onset is None:
            skipped += 1
            continue

        speaker = default_speaker
        if i_spk is not None and i_spk < len(raw):
            said = str(raw[i_spk] or "").strip()
            if said:
                if said.lower() in known:
                    speaker = known[said.lower()]
                else:
                    unknown.add(said)
                    skipped += 1
                    continue

        end = _number(raw[i_off]) if i_off is not None and i_off < len(raw) else None
        given = end is not None and end > onset
        out.append({
            "Word": word,
            "Speaker": speaker,
            "AudioFile": audio_by_speaker[speaker],
            "tStart": onset,
            "tEnd": end if given else onset + ONSET_ONLY_WINDOW,
            "Boundaries": "" if given else "FROM ONSET",
            "WordCheck": "", "EchoCheck": "", "HeardAs": "",
            "LevelDiff_dB": None,
        })

    if unknown:
        raise RuntimeError(
            "the annotation names speakers this config does not know: "
            f"{', '.join(sorted(unknown))}. The labels available are "
            f"{', '.join(audio_by_speaker)}.")
    if not out:
        raise RuntimeError(
            f"{os.path.basename(path)} gave no usable rows: every row needs a "
            "word and a numeric onset.")

    n_onset_only = sum(1 for r in out if r["Boundaries"])
    print(f"Read {len(out)} annotated word(s) from "
          f"{os.path.basename(path)} (sheet '{ws.title}').")
    if n_onset_only:
        print(f"  {n_onset_only} of them gave an onset but no end; their "
              "boundaries come from silence detection and are marked "
              '"FROM ONSET".')
    if skipped:
        print(f"  {skipped} row(s) skipped for want of a word or a numeric onset.")
    return out
